from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r'C:\code\matchalog\deliverables\MatchaLog_PRD_v1.1.docx')
d = Document()
s = d.sections[0]
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.8)
s.page_width, s.page_height = Inches(8.5), Inches(11)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st=d.styles[name]
    st.font.name='Calibri'
    st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(28)
d.styles['Heading 1'].font.size=Pt(19)
d.styles['Heading 2'].font.size=Pt(13)

def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t):
    d.add_page_break()
    d.add_heading(t,1)
def table(headers, rows, widths):
    t=d.add_table(rows=1, cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells, headers): c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=txt
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        if ri==0: pr.append(OxmlElement('w:tblHeader'))
        pr.append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=1
            cp=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8ECE9' if ri==0 else 'FFFFFF'); cp.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            cp.append(borders)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(5)
                para.paragraph_format.space_before=Pt(5)
                for r in para.runs: r.font.size=Pt(10); r.bold=ri==0
    d.add_paragraph().paragraph_format.space_after=Pt(0)

d.add_heading('MatchaLog Product Requirements',0)
p('Version 1.1 | September 5, 2026 | Product scope and release requirements')
d.add_heading('1 Product purpose and release scope',1)
p('MatchaLog helps people remember what matcha they have tried, manage what they have, and decide what to try next. The first release combines a curated product catalog, a private personal collection, and public community reviews.')
p('The product remains a responsive web application using Next.js App Router, React, JavaScript, and Supabase. Existing implementation and design conventions should be preserved unless a requirement below calls for a change.')
h('Primary users')
p('The Casual Sipper needs quick saves and a simple way to remember past purchases. The Matcha Enthusiast needs collection statuses, useful product information, and reviews that explain preparation context. Recipe creators remain a future audience rather than a separate MVP workflow.')
h('Release priorities')
table(['Release','Included scope'],[
('MVP','Curated catalog, product search and details, authentication, private stash, public reviews, basic profile/settings, and catalog suggestion and moderation tools.'),
('Next','Public profile discovery, optional collection sharing, follows, and an activity feed with explicit privacy rules.'),
('Later','Recipe creation, recipe detail, and community recipe discovery, subject to evidence of user demand.')],[1.05,5.85])
h('Changes from version 1.0')
p('Authentication and ownership controls precede multi-user writes. Stash statuses now represent ownership and use. Catalog operations, product identity, review context, privacy, product-value metrics, and testable acceptance criteria are explicit. Accessibility, responsive behavior, and error recovery apply to every phase.')
p('The previous document described Discover and basic Product Detail as built or in progress. Those are historical status notes; completion of this revision’s requirements must be verified against the implementation before release.')

page('2 Catalog and discovery')
h('Catalog ownership and missing products')
p('MVP catalog data is manually curated. Only an authorized catalog administrator may publish or edit canonical products. A signed-in user can suggest a missing product or report a correction; submission does not publish catalog content automatically.')
p('Suggestions require a product name and brand; a source link and notes are optional. A user may mark a suggestion as Want to Try. It appears in a separate pending section of their stash and becomes a normal entry when linked to an approved product. Rejection includes a reason; duplicate suggestions link to the existing product without duplicating stash entries.')
p('The curator workflow supports pending, approved, rejected, and linked-to-existing outcomes. Before launch, assign an owner for initial catalog population, suggestion triage, product corrections, and reported reviews. Define a realistic review cadence and display submission status to the requester.')
h('Product identity and fields')
p('A canonical product represents a named matcha from one brand. Package sizes share the product page and reviews. Use a separate product only when the producer markets a materially distinct product, not merely a different tin size. Harvest or batch information is optional; a new harvest alone does not create a new product in MVP.')
p('Required catalog fields are ID, name, brand, publication status, and created/updated timestamps. Origin, description, image, source link, producer-described grade, flavor notes, and harvest details are optional and must be shown as unknown when absent. Describe flavor notes as producer or curator information, separate from user reviews.')
p('Package options contain weight in grams and, when available, price amount, currency, source URL, and date checked. Prices are reference information, not live offers. Price filtering is deferred until coverage is sufficient; any future comparison must use a common weight and currency policy.')
h('Discover and product detail')
p('Discover at / searches product name and brand with the established 300 ms debounce. Provide pagination, origin and available grade filters, and sorts for newest, name, and community rating. Unknown metadata remains browsable. Preserve search/filter state on return from a product page.')
p('Cards show image or fallback, name, brand, origin when known, average rating and review count, and a stash action. Product pages at /products/[productId] show the full product data, available packages, dated prices, review summary/list, stash controls, and navigation back to discovery. Unrated products display No reviews, never a zero-star rating.')

page('3 Authentication and private collections')
h('Authentication and access')
p('Anyone may browse published products and public reviews. Saving, reviewing, suggesting products, and editing settings require a signed-in user. A login prompt preserves the intended action and returns the user to its context after authentication; the user can then complete the action.')
p('MVP supports email/password registration, login, logout, password reset, and verification handling. Google sign-in is optional after the core flow. Validate sessions on protected writes and enforce record ownership in the database. Enable row-level security before any real multi-user data is accepted. Restrict catalog writes to authorized administrators.')
p('A single-user prototype may use an isolated development environment with disposable data. World-writable production data and hardcoded production user identities are not permitted. Authentication is a prerequisite for shared stash and review work, not a later retrofit.')
h('Stash behavior')
p('The stash at /stash is private by default and in MVP. Each user has one entry per canonical product, representing current collection state rather than individual tins or lots. Multiple packages, quantities, and purchase history are outside MVP.')
table(['Status','Meaning'],[
('Want to Try','Interested in trying or buying; ownership is not implied.'),
('Unopened','Owned, but not currently opened.'),
('Open','Currently using the product.'),
('Finished','No current supply remains after use.')],[1.3,5.6])
p('A quick save defaults to Want to Try and offers a status change. Users may move directly between any statuses, including Finished back to Unopened or Open for repeat purchases. These are useful states, not a mandatory sequence.')
p('Store tried_before separately from status. Moving to Open or Finished marks it true; users can also set or correct it manually. Returning to Want to Try or Unopened preserves it. Review publication also marks an existing stash entry as tried before, but never creates a stash entry automatically.')
p('Stash supports filtering by status and tried-before value, and sorting by date added, product name, or My rating. My rating uses the owner’s review; unrated items appear last. Empty collections link to Discover and the missing-product suggestion flow.')
p('Removal deletes only the collection entry, including its local tried-before value. It does not delete the user’s review or the catalog product. Confirm this distinction in the removal interaction and provide a clear recovery path for accidental removal.')

page('4 Reviews profiles and privacy')
h('Public reviews')
p('A review requires a whole-number rating from 1 to 5. Text is optional, with a maximum of 1000 characters. Prepared as is optional, with values Straight matcha, Latte, or Cooking or baking. One editable review is allowed per user and product.')
p('No verified purchase or stash membership is required. The form asks users to review products they have tried and clearly labels publication as public. Publication does not expose private collection details. The reviewer may edit or delete their own review; the form preserves input on validation or network failure.')
p('Show the author’s public display name and avatar, rating, optional preparation context, text, and posting date; label edited reviews. Support newest, highest-rating, and lowest-rating sorts with pagination. Aggregate only currently published, non-removed reviews and update the average/count after creation, rating edits, deletion, or moderation.')
h('Profile and settings')
p('MVP /profile provides the signed-in user’s display name, avatar, optional bio, account settings, and their reviews. Display names and avatars appear publicly with reviews; email addresses never do. Other users cannot browse private stash records through profile pages or APIs.')
p('Settings explain what is public and private, support profile edits, and provide an account-deletion flow. Account deletion removes the user’s stash, reviews, profile, pending suggestions, and owned media; canonical catalog records remain. Any retained operational data and retention period must be documented before launch.')
h('Moderation and uploads')
p('Users can report inappropriate reviews. An authorized moderator can inspect reports and hide or restore a review with a recorded reason; hidden reviews are excluded from public lists and rating summaries. Profile/avatar reports use the same moderation queue. Define support ownership before release.')
p('Avatar and catalog image uploads require file type and size validation, access controls, descriptive alternatives or appropriate decorative treatment, and a visible error state. Accept JPEG, PNG, or WebP images up to 5 MB for MVP. Store media references so replaced or deleted files can be cleaned up.')
h('Collection sharing and activity privacy')
p('Private stash actions produce no public activity in MVP. Future collection sharing must be an explicit opt-in. Private records must remain inaccessible even when a visitor knows an ID or URL. Making a collection private removes its shared visibility and associated stash activity; source deletion or moderation also removes the corresponding public activity.')

page('5 Data and interface requirements')
p('This section defines required data relationships and invariants. Database migrations and detailed access policies are implementation deliverables. All user-owned writes use the authenticated identity, never a client-supplied owner ID.')
table(['Entity','Required structure and constraints'],[
('matcha_products','Canonical catalog fields from section 2; publication status; admin-controlled writes. Preserve existing IDs when curating records.'),
('product_packages','Product reference, positive weight_grams, optional nonnegative price and currency, price source and checked date. Required price metadata must accompany a displayed price.'),
('profiles','One row per auth user; display_name, avatar reference, optional bio, timestamps. Separate public fields from account/private data.'),
('user_stash','ID, user_id, product_id, status, tried_before, created_at, updated_at. Unique user/product pair; constrain statuses to the four defined values.'),
('reviews','ID, user_id, product_id, rating, optional review_text and prepared_as, visibility/moderation state, timestamps. Unique user/product pair; enforce rating and text limits.'),
('product_suggestions','Requester, proposed name/brand, optional source/notes, want-to-try intent, review status/reason, resolved product reference, timestamps. Private to requester and curator.'),
('content_reports','Reporter, target type and ID, reason, triage status, moderator outcome and timestamps. Restrict access to reporter submission and authorized moderation.')],[1.5,5.4])
h('API behavior')
p('Retain Next.js Route Handlers under app/api. GET /api/stash lists the current user’s entries; POST /api/stash saves a product idempotently; PATCH /api/stash/[id] changes state; DELETE removes the owned entry. Repeated saves must not reset an existing status.')
p('GET /api/products/[productId]/reviews lists public reviews. POST /api/reviews creates or updates the caller’s review; DELETE /api/reviews/[id] deletes only their review. Provide authenticated suggestion submission/status routes and protected curator/moderator actions. Paginate catalog, stash, and review lists with stable ordering.')
p('Validate identifiers, allowed values, text lengths, and ownership before mutation. Return understandable unauthenticated, forbidden, invalid-input, not-found, and transient-failure states. Ensure duplicate/retried requests cannot create duplicate stash or review rows. Future follows require a unique follower/following pair and prohibit self-following.')

page('6 Success measures and release acceptance')
h('Product value')
p('Collect an initial four-week baseline after beta launch, then set numerical product targets and review them monthly. Define events consistently, exclude staff/test traffic, and avoid sending review text or private stash contents to analytics.')
table(['Measure','Definition'],[
('First-session activation','Share of new signed-in users who save at least one catalog product during their first authenticated session. A session ends after 30 minutes of inactivity.'),
('Collection return rate','Share of newly activated users who return on a later calendar day within 14 days to add a product or change collection status. Count only fully observed cohorts.'),
('Repeat reviewing','Share of first-time reviewers who publish a review of a second product within 30 days. Exclude edits; use fully observed cohorts.'),
('Review completion','Successful review publications divided by sessions in which the user starts entering review data, measured weekly. Initial operational target exceeds 80%.')],[1.6,5.3])
h('Reliability and performance')
p('Target greater than 99% successful valid stash mutations, measured weekly; count failed valid requests, but exclude deliberate validation/auth rejections. Instrument catalog and product-page load performance under a documented mobile test profile and set a release budget before beta. Search-to-click time is diagnostic, not a stand-alone success target: longer comparison can reflect useful browsing.')
p('Replace the ambiguous Lighthouse mobile usability score with separate performance and accessibility checks, plus manual keyboard and mobile testing. Future feed performance receives its own budget when the social phase is scoped.')
h('Acceptance criteria for MVP')
p('Discovery: name and brand searches return matching published products; filters combine correctly; empty/error states offer recovery; pagination has stable ordering; unrated items show No reviews.')
p('Collections: repeated saves yield one entry; every status transition and repeat-purchase transition works; tried-before history survives ordinary status changes; private entries cannot be read or changed by another user; removal leaves reviews intact.')
p('Reviews: a second submission updates the existing review; invalid ratings/text are rejected; failed submissions preserve input; edits/deletion/moderation update public counts and averages; unauthorized edits fail.')
p('Catalog operations: suggestions retain requester intent, approved duplicates link to one canonical product, and rejected suggestions show a reason. Pricing is hidden when required context is missing.')
p('Usability: core flows work at narrow phone widths and with keyboard navigation; fields have labels, focus is visible, errors are announced, images have appropriate alternatives, and reduced-motion preferences are respected. Loading and retry states are present in every release flow.')

page('7 Delivery plan and future features')
h('Implementation sequence')
table(['Stage','Scope and exit condition'],[
('1 Foundation','Confirm existing behavior; establish auth, database ownership controls, profile records, catalog curation, and an initial published dataset. No production anonymous writes.'),
('2 Core MVP','Complete discovery, product details, private collections, reviews, suggestions, basic settings, and moderation. Meet section 6 acceptance criteria throughout.'),
('3 Beta validation','Instrument defined metrics, verify access boundaries and deletion behavior, test core mobile/keyboard flows, assign operational owners, and gather the baseline.'),
('4 Social follow-up','Proceed after reviewing beta usage and demand. Add profile discovery, opt-in collection sharing, follows, and a privacy-aware activity feed.'),
('5 Recipe follow-up','Proceed if demand warrants it. Scope publishing, editing, moderation, images, and discovery before estimating delivery.')],[1.25,5.65])
p('Estimate delivery after confirming implementation status, staffing, and acceptance scope. The version 1.0 week estimates are superseded; this sequence expresses dependencies rather than an unvalidated schedule.')
h('Future social requirements')
p('Profiles may expose public reviews, recipes, follow counts, and explicitly shared collections. The feed may include public review publication, shared stash additions, published recipes, and follows. Do not emit new publication events for routine edits or duplicate retries. Define whether following itself is public before launch.')
p('Each activity event must identify its source type and record, honor current source visibility, and disappear when the source is deleted or hidden. Database triggers are an implementation option, not a reason to publish private actions. Start with paginated refresh-based loading; realtime delivery requires a demonstrated need.')
h('Future recipe requirements')
p('Recipes support a title, description, ingredient amounts/units/items, ordered steps, image, optional linked product, and author. Prep time and category are required if shown on cards. Define draft versus published states, edit/delete behavior, moderation, and upload rules before implementation. Only published recipes may appear in discovery or activity.')

page('8 Technical conventions and decision register')
h('Implementation conventions')
p('Keep app/ for App Router pages and layouts, app/api/ for Route Handlers, components/ for reusable UI, and lib/ for shared utilities. Prefer Server Components for initial reads and use client components for interaction. Separate browser and server database clients as needed for session and permission handling; never expose privileged credentials to the browser.')
p('Preserve established styling: matcha green #4CAF50, light mint #edfff0, system-ui typography, and existing card/control radii where usable. Color contrast, visible focus, legibility, and touch usability take priority over literal values. Hover scaling is optional, must not shift layout, and should be disabled for reduced-motion users.')
p('Retain the established 300 ms search debounce. Use Supabase Storage for images with explicit upload ownership and cleanup behavior. Introduce SWR or React Query only when interaction and cache consistency warrant it. Verify deployed framework versions and dependencies during implementation rather than treating the original Next.js 14 label as a required version.')
h('Decisions resolved by this revision')
p('MVP product data is curated with user suggestions. Reviews require authentication and a self-reported tasting experience, but no purchase verification or stash entry. Collections are private in MVP. Authentication precedes multi-user writes. Recipes and social features are later releases. Package size is separate from canonical product identity. Collection status and prior tasting are distinct.')
h('Remaining decisions and owners')
table(['Decision','Owner and timing'],[
('Initial catalog breadth, curation cadence, and named support/moderation staff','Product / Operations, before beta.'),
('Mobile performance test profile and release budget','Engineering, before beta acceptance.'),
('Operational data retention and account deletion completion policy','Product / Engineering, before launch.'),
('Numerical activation, return, and repeat-review targets','Product, after the first four-week baseline; use mature cohorts.'),
('Public follow visibility and collection-sharing controls','Product, before the social phase.'),
('Recipe publication/moderation policy and release decision','Product, before the recipe phase.')],[3.8,3.1])
h('Out of scope for MVP')
p('Native mobile apps, commerce or vendor storefronts, AI recommendations, multilingual support, push notifications, offline mode, multiple-tin inventory and purchase history, live prices or currency conversion, social feeds/follows, and recipe publishing. These exclusions do not prevent basic responsive web access or manual catalog correction.')

OUT.parent.mkdir(parents=True,exist_ok=True)
d.core_properties.title='MatchaLog Product Requirements'
d.core_properties.subject='Revised MVP scope and release requirements'
d.core_properties.version='1.1'
for el in [d._element, d.styles.element]:
    for border in list(el.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
d.save(OUT)
print(OUT)
print('Paragraphs:',len(d.paragraphs),'Tables:',len(d.tables))
